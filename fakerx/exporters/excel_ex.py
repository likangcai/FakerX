# -*- coding: utf-8 -*-
# -----------------------------
# @Author    : 影子
# @Time      : 2026/7/14 15:13
# @Software  : PyCharm
# @FileName  : excel_ex.py
# -----------------------------
"""
Excel 导出器
"""
from typing import List, Dict

from .base import BaseExporter

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    _EXCEL_AVAILABLE = True
except ImportError:
    _EXCEL_AVAILABLE = False


class ExcelExporter(BaseExporter):
    """Excel 导出器"""

    def export(self, data: List[Dict], filepath: str, sheet_name: str = 'Sheet1',
               **kwargs):
        if not _EXCEL_AVAILABLE:
            raise ImportError("Excel 导出需要安装 openpyxl: pip install openpyxl")

        if not data:
            return

        flat_data = [self.flatten(row) for row in data]
        headers = list(flat_data[0].keys())

        wb = Workbook()
        ws = wb.active
        ws.title = sheet_name

        # 表头样式
        header_font = Font(bold=True, color='FFFFFF')
        header_fill = PatternFill(start_color='4472C4', end_color='4472C4',
                                  fill_type='solid')
        header_align = Alignment(horizontal='center', vertical='center')

        # 写入表头
        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align

        # 写入数据
        for row_idx, row in enumerate(flat_data, 2):
            for col_idx, header in enumerate(headers, 1):
                cell = ws.cell(row=row_idx, column=col_idx,
                               value=row.get(header, ''))
                cell.alignment = Alignment(vertical='top')

        # 自动列宽
        for col in ws.columns:
            max_length = 0
            column = col[0].column_letter
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except Exception:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column].width = adjusted_width

        wb.save(filepath)

    def export_stream(self, generator, filepath: str, sheet_name: str = 'Sheet1',
                      chunk_size: int = 1000, **kwargs):
        if not _EXCEL_AVAILABLE:
            raise ImportError("Excel 导出需要安装 openpyxl: pip install openpyxl")

        wb = Workbook()
        ws = wb.active
        ws.title = sheet_name

        headers_written = False
        row_idx = 2

        chunk = []
        for record in generator:
            chunk.append(record)
            if len(chunk) >= chunk_size:
                for rec in chunk:
                    flat = self.flatten(rec)
                    if not headers_written:
                        headers = list(flat.keys())
                        for col_idx, h in enumerate(headers, 1):
                            ws.cell(row=1, column=col_idx, value=h)
                        headers_written = True
                    for col_idx, h in enumerate(headers, 1):
                        ws.cell(row=row_idx, column=col_idx, value=flat.get(h, ''))
                    row_idx += 1
                chunk = []

        # 写入剩余
        for rec in chunk:
            flat = self.flatten(rec)
            if not headers_written:
                headers = list(flat.keys())
                for col_idx, h in enumerate(headers, 1):
                    ws.cell(row=1, column=col_idx, value=h)
                headers_written = True
            for col_idx, h in enumerate(headers, 1):
                ws.cell(row=row_idx, column=col_idx, value=flat.get(h, ''))
            row_idx += 1

        wb.save(filepath)
